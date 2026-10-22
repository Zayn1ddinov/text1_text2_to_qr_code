import argparse
import csv
import os
import sys

import qrcode
from PIL import Image, ImageDraw, ImageFont

# Шрифты с поддержкой кириллицы: macOS, Windows, Linux
FONT_CANDIDATES = [
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
]

TEXT_MISSING = "Надпись отсутствует"


def load_font(font_path, font_size):
    paths = [font_path] if font_path else FONT_CANDIDATES
    for path in paths:
        if path and os.path.exists(path):
            return ImageFont.truetype(path, font_size)
    if font_path:
        sys.exit(f"Шрифт не найден: {font_path}")
    # Встроенный шрифт Pillow (>= 10.1) — кириллицу может не отображать
    print("Внимание: шрифт с кириллицей не найден, используется встроенный.")
    return ImageFont.load_default(size=font_size)


# Функция для создания QR-кода с надписью
def create_qr_with_text(data, text, filename, font):
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)

    # Создание изображения QR-кода
    qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")

    # Размер надписи, чтобы выделить под неё место и выровнять по центру
    measure = ImageDraw.Draw(qr_img)
    left, top, right, bottom = measure.textbbox((0, 0), text, font=font)
    text_width, text_height = right - left, bottom - top
    padding = 10

    width = max(qr_img.size[0], text_width + 2 * padding)
    height = qr_img.size[1] + text_height + 2 * padding

    # Создание изображения с надписью
    img_with_text = Image.new("RGB", (width, height), color="white")
    img_with_text.paste(qr_img, ((width - qr_img.size[0]) // 2, 0))

    draw = ImageDraw.Draw(img_with_text)
    text_x = (width - text_width) // 2 - left
    text_y = qr_img.size[1] + padding - top
    draw.text((text_x, text_y), text, fill="black", font=font)

    # Сохранение изображения с QR-кодом и текстом
    img_with_text.save(filename)


def read_rows_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        # Excel в русской локали сохраняет CSV через «;»
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(f, dialect)]


def read_rows_xlsx(path):
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]


def read_items(path):
    """Читает пары (серийный номер, надпись): 1-я колонка — номер, 2-я — надпись."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        rows = read_rows_xlsx(path)
    else:
        rows = read_rows_csv(path)

    items = []
    for i, row in enumerate(rows):
        serial = row[0].strip() if row else ""
        if not serial:
            continue
        # Пропускаем строку заголовка
        if i == 0 and serial.lower() in ("serial", "серийный номер", "серийник", "номер"):
            continue
        text = row[1].strip() if len(row) > 1 and row[1].strip() else TEXT_MISSING
        items.append((serial, text))
    return items


def safe_filename(name):
    return "".join(c if c.isalnum() or c in "-_." else "_" for c in name)


def main():
    parser = argparse.ArgumentParser(description="Генерация QR-кодов с надписью под кодом")
    parser.add_argument("input", nargs="?", default="data.csv",
                        help="CSV или XLSX: 1-я колонка — серийный номер, 2-я — надпись (по умолчанию data.csv)")
    parser.add_argument("-o", "--output", default="qrcodes_with_text", help="папка для PNG")
    parser.add_argument("--font", help="путь к TTF-шрифту")
    parser.add_argument("--font-size", type=int, default=32, help="размер шрифта надписи")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        sys.exit(f"Файл не найден: {args.input}")

    items = read_items(args.input)
    if not items:
        sys.exit(f"В файле {args.input} нет серийных номеров")

    font = load_font(args.font, args.font_size)

    # Создание папки для сохранения QR-кодов, если ее нет
    os.makedirs(args.output, exist_ok=True)

    for serial_number, text in items:
        filename = os.path.join(args.output, f"QR_{safe_filename(serial_number)}.png")
        create_qr_with_text(serial_number, text, filename, font)

    print(f"Генерация QR-кодов завершена: {len(items)} шт. в папке {args.output}")


if __name__ == "__main__":
    main()
